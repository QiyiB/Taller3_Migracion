"""
Genera docs/Documentacion_Taller3.docx: documentación del proyecto en formato APA 7
(Times New Roman 12, interlineado doble, márgenes de 2,54 cm, sangría de primera línea,
número de página arriba a la derecha, tablas y figuras con numeración APA).

Uso:  python docs/generar_documentacion.py
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

AQUI = Path(__file__).parent
IMG = AQUI / "img"
SALIDA = AQUI / "Documentacion_Taller3.docx"
FUENTE = "Times New Roman"
CODIGO = "Courier New"
NEGRO = RGBColor(0, 0, 0)

doc = Document()
contador = {"tabla": 0, "figura": 0}


# ---------------------------------------------------------------- estilos APA
def fuente_estilo(estilo, tam=12, negrita=False, cursiva=False):
    estilo.font.name = FUENTE
    estilo.font.size = Pt(tam)
    estilo.font.bold = negrita
    estilo.font.italic = cursiva
    estilo.font.color.rgb = NEGRO
    rpr = estilo.element.get_or_add_rPr()
    fuentes = rpr.find(qn("w:rFonts"))
    if fuentes is None:
        fuentes = OxmlElement("w:rFonts")
        rpr.append(fuentes)
    for atributo in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fuentes.set(qn(atributo), FUENTE)
    # los estilos de título heredan la fuente del tema (Calibri Light); se elimina esa herencia
    for atributo in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if fuentes.get(qn(atributo)) is not None:
            del fuentes.attrib[qn(atributo)]


normal = doc.styles["Normal"]
fuente_estilo(normal)
normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(0)

for nivel, alineacion, cursiva in [(1, WD_ALIGN_PARAGRAPH.CENTER, False), (2, WD_ALIGN_PARAGRAPH.LEFT, False),
                                   (3, WD_ALIGN_PARAGRAPH.LEFT, True)]:
    estilo = doc.styles[f"Heading {nivel}"]
    fuente_estilo(estilo, negrita=True, cursiva=cursiva)
    pf = estilo.paragraph_format
    pf.alignment = alineacion
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Cm(0)
    pf.keep_with_next = True

for nombre in ("List Bullet", "List Number"):
    estilo = doc.styles[nombre]
    fuente_estilo(estilo)
    estilo.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

seccion = doc.sections[0]
seccion.page_width, seccion.page_height = Inches(8.5), Inches(11)
for lado in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(seccion, lado, Inches(1))

# número de página arriba a la derecha (APA 7, trabajo de estudiante)
encabezado = seccion.header.paragraphs[0]
encabezado.alignment = WD_ALIGN_PARAGRAPH.RIGHT
corrida = encabezado.add_run()
for tipo, texto_campo in (("begin", None), (None, "PAGE"), ("end", None)):
    if tipo:
        campo = OxmlElement("w:fldChar")
        campo.set(qn("w:fldCharType"), tipo)
    else:
        campo = OxmlElement("w:instrText")
        campo.set(qn("xml:space"), "preserve")
        campo.text = texto_campo
    corrida._r.append(campo)
corrida.font.name = FUENTE
corrida.font.size = Pt(12)


# ---------------------------------------------------------------- utilidades
def parrafo(texto="", sangria=True, alineacion=None, negrita=False, cursiva=False):
    """texto: str o lista de (texto, {'b': bool, 'i': bool, 'code': bool})."""
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0.5) if sangria else Cm(0)
    if alineacion is not None:
        p.alignment = alineacion
    partes = texto if isinstance(texto, list) else [(texto, {})]
    for t, o in partes:
        r = p.add_run(t)
        r.bold = o.get("b", negrita)
        r.italic = o.get("i", cursiva)
        if o.get("code"):
            r.font.name = CODIGO
            r.font.size = Pt(11)
    return p


def titulo(texto, nivel):
    return doc.add_heading(texto, level=nivel)


def titulo4(etiqueta, texto):
    """Nivel 4 de APA: con sangría, negrita, terminado en punto y seguido del texto."""
    return parrafo([(etiqueta + ". ", {"b": True})] + (texto if isinstance(texto, list) else [(texto, {})]))


def vinetas(items, numerada=False):
    for item in items:
        p = doc.add_paragraph(style="List Number" if numerada else "List Bullet")
        partes = item if isinstance(item, list) else [(item, {})]
        for t, o in partes:
            r = p.add_run(t)
            r.bold = o.get("b", False)
            r.italic = o.get("i", False)
            if o.get("code"):
                r.font.name = CODIGO
                r.font.size = Pt(11)


def codigo(lineas):
    """Bloque de código: Courier New 10, interlineado sencillo, con sangría."""
    for i, linea in enumerate(lineas):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.left_indent = Inches(0.25)
        pf.first_line_indent = Cm(0)
        pf.space_after = Pt(12) if i == len(lineas) - 1 else Pt(0)
        pf.keep_with_next = i < len(lineas) - 1
        r = p.add_run(linea)
        r.font.name = CODIGO
        r.font.size = Pt(10)


def borde(celda, lado, valor="single", tam=8):
    tc_pr = celda._tc.get_or_add_tcPr()
    bordes = tc_pr.find(qn("w:tcBorders"))
    if bordes is None:
        bordes = OxmlElement("w:tcBorders")
        tc_pr.append(bordes)
    b = OxmlElement(f"w:{lado}")
    b.set(qn("w:val"), valor)
    b.set(qn("w:sz"), str(tam))
    b.set(qn("w:color"), "000000")
    bordes.append(b)


def etiqueta_numerada(tipo, titulo_texto):
    contador[tipo] += 1
    nombre = "Tabla" if tipo == "tabla" else "Figura"
    p = parrafo(f"{nombre} {contador[tipo]}", sangria=False, negrita=True)
    p.paragraph_format.keep_with_next = True
    p = parrafo(titulo_texto, sangria=False, cursiva=True)
    p.paragraph_format.keep_with_next = True
    return contador[tipo]


def nota(texto):
    p = parrafo([("Nota. ", {"i": True}), (texto, {})], sangria=False)
    p.paragraph_format.space_after = Pt(12)


def tabla(titulo_texto, encabezados, filas, anchos, nota_texto=None, alinear_derecha=()):
    """Tabla APA: solo líneas horizontales (arriba, bajo el encabezado y al final)."""
    etiqueta_numerada("tabla", titulo_texto)
    t = doc.add_table(rows=1 + len(filas), cols=len(encabezados))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for fila_i, valores in enumerate([encabezados] + filas):
        fila = t.rows[fila_i]
        for col_i, valor in enumerate(valores):
            celda = fila.cells[col_i]
            celda.width = Inches(anchos[col_i])
            celda.text = ""
            p = celda.paragraphs[0]
            pf = p.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
            pf.first_line_indent = Cm(0)
            pf.space_before = Pt(3)
            pf.space_after = Pt(3)
            pf.keep_with_next = fila_i < len(filas)    # la tabla no se separa de su encabezado
            if col_i in alinear_derecha and fila_i > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            partes = valor if isinstance(valor, list) else [(str(valor), {})]
            for txt, o in partes:
                r = p.add_run(txt)
                r.bold = fila_i == 0 or o.get("b", False)
                r.italic = o.get("i", False)
                if o.get("code"):
                    r.font.name = CODIGO
                    r.font.size = Pt(10)
            if fila_i == 0:
                borde(celda, "top")
                borde(celda, "bottom")
            if fila_i == len(filas):
                borde(celda, "bottom")
        tr_pr = fila._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))       # una fila no se parte entre páginas
        if fila_i == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))   # el encabezado se repite en cada página
    if nota_texto:
        nota(nota_texto)
    else:
        doc.add_paragraph().paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return t


def figura(titulo_texto, archivo, ancho=6.5, nota_texto=None):
    etiqueta_numerada("figura", titulo_texto)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(archivo), width=Inches(ancho))
    if nota_texto:
        nota(nota_texto)


def salto_pagina():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


C = {"code": True}
B = {"b": True}
I = {"i": True}

# ================================================================= portada
for _ in range(3):
    parrafo("", sangria=False)
parrafo("Migración de una Base de Datos Relacional de Comercio Electrónico desde SQL Server hacia MongoDB",
        sangria=False, alineacion=WD_ALIGN_PARAGRAPH.CENTER, negrita=True)
parrafo("Documentación del Taller Evaluativo No. 3", sangria=False, alineacion=WD_ALIGN_PARAGRAPH.CENTER)
parrafo("", sangria=False)
for linea in ["[Nombre del integrante 1] y [Nombre del integrante 2]",
              "Universidad Católica Luis Amigó",
              "Bases de Datos II",
              "Profesor: Juan Carlos Gómez",
              "22 de octubre de 2026"]:
    parrafo(linea, sangria=False, alineacion=WD_ALIGN_PARAGRAPH.CENTER)
salto_pagina()

# ================================================================= título e introducción
parrafo("Migración de una Base de Datos Relacional de Comercio Electrónico desde SQL Server hacia MongoDB",
        sangria=False, alineacion=WD_ALIGN_PARAGRAPH.CENTER, negrita=True)
parrafo("Este documento describe el proyecto desarrollado para el Taller Evaluativo No. 3 del curso Bases de "
        "Datos II. El trabajo parte de una base de datos relacional construida en SQL Server para un caso de "
        "comercio electrónico, la analiza y la transforma en un modelo documental en MongoDB. La "
        "transformación no copia cada tabla como una colección: decide, para cada relación, si la información "
        "debe almacenarse dentro de un mismo documento (Embedding) o en colecciones independientes "
        "relacionadas por identificadores (Referencing), y justifica cada decisión.")
parrafo("El documento sigue el orden de los entregables solicitados: los scripts de creación y carga en SQL "
        "Server, el código de migración, el modelo relacional y el modelo documental final, las evidencias de "
        "la estructura y de los datos migrados, las verificaciones posteriores a la migración, la explicación "
        "de las decisiones de Embedding y Referencing, la justificación de las consultas para el negocio y la "
        "presentación de la sustentación.")

titulo("Contenido del Repositorio", 2)
parrafo("La Tabla 1 relaciona cada entregable solicitado con su ubicación dentro del repositorio Git del "
        "proyecto. Todas las rutas son relativas a la carpeta raíz del repositorio.")
tabla("Entregables del Taller y su Ubicación en el Repositorio",
      ["Entregable", "Ubicación en el repositorio"],
      [["Scripts de creación y carga de la base de datos en SQL Server",
        [("sql/01_esquema.sql", C), (", ", {}), ("sql/02_carga_datos.sql", C), (", ", {}),
         ("sql/03_verificacion_sql.sql", C), (" y ", {}), ("sql/generar_datos.py", C)]],
       ["Código utilizado para la migración hacia MongoDB", [("migracion/migrar.py", C), (" y ", {}),
                                                             ("config.py", C)]],
       ["Código de las consultas y Aggregation Pipelines", [("consultas/consultas.js", C)]],
       ["Modelo relacional y modelo documental final", [("docs/modelo_relacional.md", C), (" y ", {}),
                                                        ("docs/modelo_documental.md", C)]],
       ["Evidencias de la estructura y los datos migrados", [("evidencias/ejemplos/", C), (" y ", {}),
                                                             ("evidencias/02_log_migracion.txt", C)]],
       ["Evidencias de las verificaciones", [("evidencias/01_verificacion_sql_server.txt", C), (", ", {}),
                                             ("evidencias/03_verificacion_migracion.md", C), (" y ", {}),
                                             ("migracion/verificar.py", C)]],
       ["Explicación de Embedding y Referencing", [("docs/modelo_documental.md", C),
                                                   (" y este documento", {})]],
       ["Justificación de las consultas", [("docs/consultas.md", C), (" y ", {}),
                                           ("evidencias/04_resultados_consultas.txt", C)]],
       ["Presentación de la sustentación", [("presentacion/Sustentacion_Taller3.pptx", C)]]],
      [2.6, 3.9],
      "El repositorio incluye además una aplicación de escritorio complementaria (carpeta migrador) que se "
      "describe al final de este documento.")

# ================================================================= caso de estudio
titulo("Caso de Estudio", 1)
parrafo("El caso de estudio es TiendaOnline, una tienda de comercio electrónico colombiana que vende 300 "
        "productos de 12 categorías a 800 clientes ubicados en 12 ciudades del país. Los pedidos se pagan "
        "con cinco medios (tarjeta de crédito, tarjeta débito, PSE, Nequi y pago contraentrega) y se "
        "despachan con seis transportadoras. Los clientes pueden calificar los productos que recibieron. Los "
        "datos cubren el periodo comprendido entre enero de 2025 y septiembre de 2026.")
parrafo("Se eligió este dominio porque combina entidades maestras (productos, clientes, proveedores y "
        "categorías) con entidades transaccionales (pedidos, líneas de pedido, pagos y envíos). Esa "
        "combinación obliga a tomar decisiones reales de modelado documental, ya que unas entidades tienen "
        "existencia propia y se consultan desde muchas partes del sistema, mientras que otras solo tienen "
        "sentido dentro del pedido al que pertenecen.")

# ================================================================= modelo relacional
titulo("Modelo Relacional en SQL Server", 1)
parrafo("La base de datos TiendaOnline se construyó en SQL Server 2025 Express y está compuesta por 10 "
        "entidades relacionadas mediante 11 llaves foráneas. Contiene 14.027 registros distribuidos según la "
        "lógica del negocio: las tablas maestras tienen pocos registros y las transaccionales concentran el "
        "mayor volumen, como se observa en la Tabla 2.")
tabla("Entidades del Modelo Relacional y Volumen de Registros",
      ["Tabla", "Registros", "Rol en el negocio"],
      [["Categorias", "12", "Clasificación del catálogo"],
       ["Proveedores", "25", "Empresas que abastecen los productos"],
       ["Productos", "300", "Catálogo con precio, costo y existencias"],
       ["Clientes", "800", "Compradores registrados y su segmento (Estándar, Plata u Oro)"],
       ["Direcciones", "1.185", "Direcciones de entrega de cada cliente (de una a tres)"],
       ["Pedidos", "2.000", "Encabezado de cada compra y su estado"],
       ["DetallePedido", "4.317", "Líneas de cada pedido; resuelve la relación muchos a muchos entre "
                                  "pedidos y productos"],
       ["Pagos", "2.144", "Intentos de pago aprobados, rechazados o reembolsados"],
       ["Envios", "1.866", "Despacho del pedido con transportadora y fechas"],
       ["Resenas", "1.378", "Calificación de 1 a 5 de un producto recibido"],
       [[("Total", B)], [("14.027", B)], ""]],
      [1.5, 1.0, 4.0], alinear_derecha=(1,))

titulo("Relaciones", 2)
parrafo("La Figura 1 presenta el diagrama entidad relación. Todas las relaciones son de uno a muchos, excepto "
        "la relación entre Pedidos y Envios, que es de uno a cero o uno porque un pedido puede no haberse "
        "despachado todavía. Cada pedido se relaciona con una de las direcciones de su propio cliente.")
figura("Diagrama Entidad Relación de TiendaOnline", IMG / "figura_modelo_relacional.png", 6.5,
       "Cada caja muestra la tabla y su número de registros; las etiquetas sobre las líneas indican la "
       "cardinalidad de la relación.")

titulo("Reglas de Integridad", 2)
parrafo("Además de las llaves primarias y foráneas, el esquema define reglas que impiden datos incoherentes:")
vinetas([
    [("Restricciones CHECK: ", B), ("precios y costos positivos, existencias no negativas, estados y medios "
                                   "de pago dentro de un dominio cerrado, calificación entre 1 y 5, descuento "
                                   "entre 0 y 100, y la regla Total = Subtotal + CostoEnvio.", {})],
    [("Restricciones UNIQUE: ", B), ("SKU del producto, correo del cliente, NIT del proveedor, referencia de "
                                    "pago, número de guía, un mismo producto una sola vez por pedido, una "
                                    "reseña por cliente y producto, y un envío como máximo por pedido.", {})],
    [("Columna calculada persistida: ", B), ("el subtotal de cada línea de pedido se calcula como cantidad "
                                            "por precio unitario menos el descuento.", {})],
    [("Índices: ", B), ("se crearon índices sobre todas las llaves foráneas, porque SQL Server no los crea "
                       "automáticamente.", {})],
])

titulo("Scripts de Creación y Carga", 2)
parrafo([("El script ", {}), ("sql/01_esquema.sql", C), (" crea la base de datos y las 10 tablas con sus "
         "restricciones e índices. El script ", {}), ("sql/02_carga_datos.sql", C),
         (" inserta los 14.027 registros. Este segundo script lo genera el programa ", {}),
         ("sql/generar_datos.py", C), (", que produce datos sintéticos con una semilla fija, de modo que el "
         "resultado siempre es el mismo.", {})])
parrafo("La generación respeta reglas de negocio para que los datos sean coherentes: cada pedido se envía a una "
        "dirección de su propio cliente, el subtotal del pedido es la suma de sus líneas, el envío es gratuito "
        "desde 150.000 pesos, el estado del pedido corresponde a sus pagos y a su envío, el pago contraentrega "
        "solo aparece en pedidos entregados y solo reseñan un producto los clientes que lo recibieron. "
        "Además, cada transportadora tiene un tiempo de entrega distinto, cada medio de pago una tasa de "
        "rechazo distinta y cada producto una calidad que influye en sus calificaciones, lo que permite que "
        "las consultas encuentren diferencias reales.")
parrafo([("El script ", {}), ("sql/03_verificacion_sql.sql", C),
         (" comprueba el volumen por tabla y cinco controles de coherencia; los cinco dieron cero casos "
          "(evidencia en ", {}), ("evidencias/01_verificacion_sql_server.txt", C), (").", {})])

# ================================================================= migración
titulo("Proceso de Migración hacia MongoDB", 1)
parrafo([("La migración está implementada en Python en el archivo ", {}), ("migracion/migrar.py", C),
         (", con las bibliotecas pyodbc para leer SQL Server y pymongo para escribir en MongoDB. Los "
          "parámetros de conexión se encuentran en ", {}), ("config.py", C),
         (" y pueden cambiarse mediante variables de entorno. El proceso tiene tres etapas.", {})])
titulo4("Extracción", "Se leen las 10 tablas con consultas SELECT ordenadas por su llave. Las líneas de pedido "
        "se leen con un JOIN a Productos para obtener el SKU y el nombre de cada producto.")
titulo4("Transformación", "Las filas hijas se agrupan por su llave foránea con un diccionario "
        "(defaultdict). Así, cada cliente recibe sus direcciones y cada pedido recibe sus líneas, sus pagos y "
        "su envío en un solo paso. La dirección del pedido se busca por su identificador y se copia dentro del "
        "pedido. Las reseñas se agrupan por producto para calcular el promedio y el total de calificaciones.")
titulo4("Carga", "Cada colección se elimina y se vuelve a crear, de modo que la migración se puede ejecutar "
        "varias veces sin duplicar datos. Se aplican los validadores, los documentos se insertan en lotes de "
        "1.000 y se crean los índices. La migración completa tarda alrededor de un segundo.")

titulo("Decisiones Técnicas de la Migración", 2)
vinetas([
    [("Identificadores: ", B), ("el campo _id de cada documento conserva la llave primaria de SQL Server. "
                                "Esto mantiene la trazabilidad (el pedido 24 es el mismo en ambos motores) "
                                "y simplifica la verificación.", {})],
    [("Valores monetarios: ", B), ("los campos DECIMAL se convierten a Decimal128, que conserva la precisión "
                                   "exacta. Con números de punto flotante aparecerían errores de redondeo en "
                                   "las sumas.", {})],
    [("Fechas: ", B), ("los campos DATE se convierten a fecha y hora a medianoche, porque MongoDB no tiene un "
                       "tipo que almacene solo la fecha.", {})],
    [("Valores nulos: ", B), ("un campo que no aplica no se guarda. Por ejemplo, un pedido sin despacho no "
                              "tiene el campo envio.", {})],
    [("Validación: ", B), ("las colecciones pedidos y clientes tienen un validador $jsonSchema con campos "
                           "obligatorios, estados permitidos y al menos una línea por pedido.", {})],
    [("Índices: ", B), ("se crearon índices sobre las referencias y los campos de filtro más usados, entre "
                        "ellos cliente y fecha del pedido, producto dentro de las líneas, transportadora, "
                        "categoría, proveedor, existencias, SKU y correo.", {})],
])

# ================================================================= modelo documental
titulo("Modelo Documental Final", 1)
parrafo("Las 10 tablas relacionales se transformaron en seis colecciones. La Tabla 3 resume el origen de cada "
        "colección y la técnica aplicada, y la Figura 2 muestra el destino de cada tabla.")
tabla("Colecciones del Modelo Documental",
      ["Colección", "Documentos", "Proviene de", "Técnica aplicada"],
      [[[("categorias", C)], "12", "Categorias", "Colección independiente"],
       [[("proveedores", C)], "25", "Proveedores", "Colección independiente; contacto y ubicación agrupados "
                                                  "en subdocumentos"],
       [[("productos", C)], "300", "Productos y Resenas", "Referencia a categoría y proveedor; promedio de "
                                                         "calificación calculado"],
       [[("clientes", C)], "800", "Clientes y Direcciones", "Direcciones embebidas como array"],
       [[("pedidos", C)], "2.000", "Pedidos, DetallePedido, Pagos, Envios y Direcciones",
        "Líneas, pagos y envío embebidos; copia de la dirección; referencia al cliente y a los productos"],
       [[("resenas", C)], "1.378", "Resenas", "Referencia al producto y al cliente"]],
      [1.25, 1.15, 1.8, 2.3], alinear_derecha=(1,))
figura("Destino de Cada Tabla Relacional en el Modelo Documental", IMG / "figura_mapeo_colecciones.png", 6.5,
       "Las líneas grises indican tablas que conservan su propia colección; las naranjas, tablas que se "
       "embeben dentro de otra colección; las punteadas, datos copiados o calculados.")

titulo("Estructura del Documento de Pedido", 2)
parrafo("La colección pedidos reúne cuatro tablas en un solo documento. El siguiente fragmento, tomado del "
        "pedido 24 real, muestra su estructura: el cliente y los productos se referencian por su "
        "identificador, mientras que las líneas, los pagos y el envío quedan embebidos.")
codigo([
    "{",
    "  _id: 24,",
    "  cliente_id: 658,                   // referencia a clientes",
    "  fecha: ISODate(\"2025-01-15T23:56:54Z\"),",
    "  estado: \"Entregado\",",
    "  direccion_envio: {                 // copia de la dirección",
    "    direccion_id: 985, ciudad: \"Medellín\", ... },",
    "  items: [                           // antes DetallePedido",
    "    { producto_id: 69, sku: \"COC-00069\",",
    "      nombre: \"Sartén antiadherente Kora Pro\",",
    "      cantidad: 1,",
    "      precio_unitario: Decimal128(\"174900.00\"),",
    "      descuento_pct: Decimal128(\"0.00\"),",
    "      subtotal: Decimal128(\"174900.00\") },",
    "    ...",
    "  ],",
    "  totales: { subtotal: ..., costo_envio: ...,",
    "             total: Decimal128(\"585220.00\") },",
    "  pagos: [                           // antes Pagos",
    "    { pago_id: 25, metodo: \"Tarjeta credito\",",
    "      estado: \"Rechazado\", ... },",
    "    { pago_id: 26, metodo: \"Tarjeta credito\",",
    "      estado: \"Aprobado\", ... }",
    "  ],",
    "  envio: {                           // antes Envios",
    "    transportadora: \"Servientrega\", numero_guia: \"SVT002408767\",",
    "    fecha_entrega_estimada: ..., fecha_entrega: ...,",
    "    estado: \"Entregado\" }",
    "}",
])
parrafo([("Las colecciones clientes y productos siguen la misma lógica. Cada cliente contiene el array ", {}),
         ("direcciones", C), (" con sus direcciones, y cada producto conserva las referencias ", {}),
         ("categoria_id", C), (" y ", {}), ("proveedor_id", C), (" y el subdocumento calculado ", {}),
         ("calificacion", C), (" con el promedio y el total de reseñas. Los documentos reales de cada "
          "colección se encuentran en ", {}), ("evidencias/ejemplos/", C), (".", {})])

# ================================================================= embedding vs referencing
titulo("Decisiones de Embedding y Referencing", 1)
parrafo("El criterio general fue el siguiente: se embebe la información que pertenece a un solo padre, que "
        "tiene un tamaño acotado y que casi siempre se lee junto con ese padre; se referencia la información "
        "que tiene existencia propia, que se usa desde varias partes del sistema o que puede crecer sin "
        "límite. La Tabla 4 aplica este criterio a cada relación del modelo.")
tabla("Decisión de Modelado para Cada Relación",
      ["Relación", "Decisión", "Justificación"],
      [["Pedidos y DetallePedido", [("Embedding (", {}), ("items", C), (")", {})],
        "Una línea no existe sin su pedido y nunca se consulta sola. Un pedido tiene de una a cinco líneas, "
        "por lo que el array está acotado. Crear o cancelar un pedido es una escritura atómica sobre un solo "
        "documento."],
       ["Pedidos y Pagos", [("Embedding (", {}), ("pagos", C), (")", {})],
        "Los intentos de pago pertenecen solo a su pedido y son pocos (normalmente uno, y como máximo dos). "
        "Para saber si un pedido está pagado hay que verlos junto a él."],
       ["Pedidos y Envios", [("Embedding (", {}), ("envio", C), (")", {})],
        "Es una relación de uno a cero o uno. Una colección aparte para un único documento por pedido solo "
        "agregaría operaciones $lookup."],
       ["Pedidos y Direcciones", [("Copia (", {}), ("direccion_envio", C), (")", {})],
        "Si el cliente cambia o elimina la dirección, el pedido debe seguir mostrando a dónde se envió "
        "realmente. Es un dato histórico, como en una factura."],
       ["Clientes y Direcciones", [("Embedding (", {}), ("direcciones", C), (")", {})],
        "Cada cliente tiene de una a tres direcciones que solo le pertenecen a él y que se muestran en su "
        "perfil y al comprar."],
       ["Pedidos y Clientes", [("Referencing (", {}), ("cliente_id", C), (")", {})],
        "El cliente tiene existencia propia y acumula pedidos sin límite. Embeber los pedidos en el cliente "
        "haría crecer el documento hacia el límite de 16 MB, y embeber el cliente en cada pedido duplicaría "
        "sus datos."],
       ["Líneas y Productos", [("Referencing (", {}), ("producto_id", C), (") con copia de SKU y nombre", {})],
        "El producto lo usan el catálogo, el inventario, los pedidos y las reseñas, y su precio y "
        "existencias cambian. El precio de la línea se guarda porque es el precio al que se vendió."],
       ["Productos con Categorias y Proveedores", "Referencing",
        "Son catálogos compartidos por muchos productos; un cambio se hace en un solo documento."],
       ["Resenas", "Colección propia con referencias",
        "Las reseñas crecen sin límite, se consultan desde el producto y desde el cliente, y se analizan por "
        "separado."],
       ["Productos y Resenas", [("Campo calculado (", {}), ("calificacion", C), (")", {})],
        "La ficha del producto muestra su promedio en cada visita; precalcularlo evita agregar todas las "
        "reseñas en cada lectura."]],
      [1.55, 1.6, 3.35])

titulo("Ventajas para el Caso", 2)
vinetas([
    "La consulta de un pedido completo se resuelve con una sola lectura por identificador, mientras que en "
    "SQL Server requiere cinco operaciones JOIN.",
    "Crear un pedido con sus líneas y su primer pago es una sola inserción atómica.",
    "Las entidades que cambian (cliente, producto, proveedor y categoría) existen una sola vez y se combinan "
    "con $lookup cuando hace falta.",
    "El precio y la dirección del pedido quedan registrados tal como estaban en el momento de la compra.",
])
titulo("Costos Asumidos", 2)
parrafo("La copia del nombre del producto dentro de las líneas puede quedar desactualizada si el producto se "
        "renombra. Se acepta porque refleja cómo se llamaba el producto cuando se vendió. Por otra parte, el "
        "promedio de calificación debe recalcularse cuando llega una reseña nueva; en la migración se calcula "
        "y la verificación comprueba que coincide con las reseñas almacenadas.")

# ================================================================= verificación
titulo("Verificación de la Migración", 1)
parrafo([("El programa ", {}), ("migracion/verificar.py", C), (" compara SQL Server con MongoDB después de "
         "cada migración. Se realizaron 79 controles y los 79 fueron correctos. La Tabla 5 resume los tipos "
         "de control.", {})])
tabla("Controles de Verificación Aplicados",
      ["Tipo de control", "Cantidad", "Ejemplo"],
      [["Conteos", "15", "Las 4.317 filas de DetallePedido equivalen a la suma de los tamaños del array items "
                         "(4.317). También se comparan los pedidos por estado."],
       ["Sumas de control", "6", "La suma del total de los pedidos es 1.426.165.610 pesos en ambos motores, "
                                 "sin diferencia alguna."],
       ["Integridad referencial", "6", "Cada referencia se resolvió con $lookup y no se encontró ninguna "
                                       "referencia sin destino."],
       ["Estructuras embebidas", "7", "En los 2.000 pedidos, la suma de los subtotales de las líneas es igual "
                                      "al subtotal del pedido."],
       ["Muestra campo a campo", "45", "Cinco pedidos elegidos al azar (309, 664, 809, 1.334 y 1.942) "
                                       "comparados campo por campo y línea por línea."],
       [[("Total", B)], [("79", B)], "79 de 79 controles correctos"]],
      [1.8, 0.9, 3.8], "El reporte completo se encuentra en evidencias/03_verificacion_migracion.md.",
      alinear_derecha=(1,))

titulo("Verificación de Tablas Fusionadas", 2)
parrafo("Cuando varias tablas se transforman en una sola estructura documental, el conteo directo de "
        "documentos no basta. Por eso las filas de cada tabla embebida se compararon con la suma de los "
        "tamaños del array correspondiente: Direcciones con el array direcciones de los clientes, "
        "DetallePedido con items y Pagos con pagos. Los envíos se compararon con la cantidad de pedidos que "
        "tienen el campo envio. Las sumas de control confirman que no solo se conservó la cantidad de "
        "elementos, sino también sus valores.")

titulo("Evidencias de la Estructura y los Datos Migrados", 2)
parrafo([("La carpeta ", {}), ("evidencias", C), (" contiene el registro de la migración (", {}),
         ("02_log_migracion.txt", C), ("), un documento real de cada colección en ", {}),
         ("evidencias/ejemplos/", C), (" y la lista de colecciones con su cantidad de documentos e índices, "
          "resumida en la Tabla 6.", {})])
tabla("Colecciones, Documentos e Índices en MongoDB",
      ["Colección", "Documentos", "Índices"],
      [[[("categorias", C)], "12", [("_id", C)]],
       [[("clientes", C)], "800", [("_id, email (único)", C)]],
       [[("pedidos", C)], "2.000", [("_id, cliente_id + fecha, fecha, estado, items.producto_id, "
                                     "envio.transportadora", C)]],
       [[("productos", C)], "300", [("_id, sku (único), categoria_id, proveedor_id, stock", C)]],
       [[("proveedores", C)], "25", [("_id", C)]],
       [[("resenas", C)], "1.378", [("_id, producto_id + cliente_id (único), cliente_id", C)]]],
      [1.3, 1.1, 4.1], alinear_derecha=(1,))

# ================================================================= consultas
titulo("Consultas y su Importancia para el Negocio", 1)
parrafo([("Se desarrollaron siete consultas con Aggregation Pipeline, ubicadas en ", {}),
         ("consultas/consultas.js", C), (". Cada una responde una pregunta concreta de la operación y apoya "
          "una decisión. La Tabla 7 las resume y las secciones siguientes explican qué información busca cada "
          "una, por qué es importante, qué decisión apoya y qué impacto tiene. Los resultados completos están "
          "en ", {}), ("evidencias/04_resultados_consultas.txt", C), (".", {})])
tabla("Resumen de las Consultas",
      ["N.º", "Pregunta de negocio", "Colección", "$lookup"],
      [["C1", "¿Qué productos sostienen la facturación?", [("pedidos", C)], "3"],
       ["C2", "¿Qué categorías dejan más margen?", [("pedidos", C)], "2"],
       ["C3", "¿Qué clientes merecen subir de segmento?", [("pedidos", C)], "1"],
       ["C4", "¿Qué transportadora cumple mejor?", [("pedidos", C)], "Ninguno"],
       ["C5", "¿Qué productos están mal calificados y de qué proveedor son?", [("resenas", C)], "4"],
       ["C6", "¿Qué productos se van a agotar pronto?", [("productos", C)], "2"],
       ["C7", "¿Qué medio de pago pierde más ventas?", [("pedidos", C)], "Ninguno"]],
      [0.55, 3.75, 1.2, 1.0],
      "C4 y C7 no requieren $lookup porque el envío y los pagos están embebidos en el pedido.")

consultas = [
    ("C1. Productos que Sostienen la Facturación",
     "Los diez productos con más ingresos en 2026, con su categoría, su proveedor, las unidades vendidas, el "
     "margen bruto y las existencias actuales.",
     "Filtra los pedidos con pago aprobado de 2026 ($match), separa cada línea ($unwind), agrupa por producto "
     "($group), ordena y limita a diez antes de los $lookup para buscar solo diez productos, y luego aplica "
     "tres $lookup encadenados (productos, categorías y proveedores) y un $project que calcula el margen como "
     "ingresos menos unidades por costo.",
     "Muestra qué productos sostienen la facturación y si además dejan margen.",
     "Asegurar el inventario de los productos estrella y negociar volumen con su proveedor.",
     "El Smartwatch Zenit Plus es el producto con más ingresos (44,1 millones de pesos y 47 % de margen), "
     "pero solo le quedan nueve unidades.",
     "Evita perder ventas en los productos que más dinero aportan."),
    ("C2. Rentabilidad por Categoría",
     "Pedidos, unidades, ingresos, margen bruto y porcentaje de margen de cada categoría.",
     "Separa las líneas, aplica un $lookup a productos con un sub-pipeline que trae solo la categoría y el "
     "costo, agrupa por categoría usando $addToSet para contar pedidos distintos, aplica un $lookup a "
     "categorías y calcula el porcentaje de margen.",
     "Las categorías que más facturan no siempre son las más rentables.",
     "Dónde invertir el presupuesto de publicidad y qué categorías ampliar.",
     "Deportes factura más que cualquier otra categoría (269 millones de pesos), pero su margen (38 %) es "
     "inferior al de Tecnología (42,8 %) y al de Oficina (42,4 %).",
     "El presupuesto de mercadeo se dirige a donde cada peso vendido deja más utilidad."),
    ("C3. Clientes Candidatos a Subir de Segmento",
     "Los 15 clientes que no pertenecen al segmento Oro y que más han comprado, con su ciudad principal, "
     "número de pedidos, compra promedio y fecha de la última compra.",
     "Agrupa los pedidos entregados por cliente, aplica un $lookup a clientes, filtra con $match un campo "
     "traído por el $lookup (segmento distinto de Oro) y obtiene la ciudad principal del array embebido de "
     "direcciones con $indexOfArray y $arrayElemAt.",
     "Si la segmentación está desactualizada, los mejores clientes no reciben los beneficios que les "
     "corresponden y pueden irse a la competencia.",
     "Reclasificar a estos clientes y dirigirles una campaña de fidelización.",
     "Hay clientes del segmento Estándar con 28 pedidos y 23,9 millones de pesos en compras.",
     "Mayor retención de los clientes de más valor."),
    ("C4. Desempeño de las Transportadoras",
     "Por transportadora: envíos entregados, días promedio y máximos de entrega, porcentaje de entregas "
     "tarde y valor transportado.",
     "Filtra por un campo del subdocumento embebido envio sin necesidad de $lookup, calcula los días reales "
     "con $dateDiff, marca con $cond las entregas que llegaron después del día prometido y agrupa por "
     "transportadora.",
     "La entrega tarde es una de las principales causas de reclamos en el comercio electrónico y afecta la "
     "recompra.",
     "A qué transportadora asignar más volumen y a cuál exigirle mejoras.",
     "Deprisa promete dos días y llega tarde en el 61,5 % de los envíos; Servientrega mueve el mayor valor "
     "y es la más cumplida, con 28,5 % de entregas tarde.",
     "Menos reclamos y una mejor experiencia de compra."),
    ("C5. Productos Mal Calificados y su Proveedor",
     "Los productos con promedio inferior a 3 y al menos tres reseñas, con su proveedor, el porcentaje de "
     "reseñas negativas y las dos quejas más recientes.",
     "Agrupa las reseñas por producto, filtra por promedio y cantidad, y aplica cuatro $lookup: productos, "
     "categorías, proveedores y un $lookup con let y pipeline sobre la misma colección de reseñas que trae "
     "las dos quejas más recientes.",
     "Un producto mal calificado genera devoluciones, reclamos y daña la reputación de la tienda.",
     "Abrir un reclamo de calidad con el proveedor, retirar el producto o cambiar de proveedor.",
     "Importadora Global S.A.S. provee tres de los productos peor calificados; la Licuadora Cumbre Plus "
     "(promedio de 1,69 en 13 reseñas) también aparece entre los diez productos con más ingresos.",
     "Menos devoluciones y una base objetiva para negociar con los proveedores."),
    ("C6. Riesgo de Agotamiento de Existencias",
     "Los productos activos con 20 unidades o menos que sí se vendieron en los últimos 90 días, con los días "
     "de inventario restantes al ritmo actual y el contacto del proveedor.",
     "Filtra los productos, aplica un $lookup a pedidos con let y un sub-pipeline que suma las unidades "
     "vendidas del producto en 90 días, descarta con $unwind los productos sin ventas, aplica un $lookup a "
     "proveedores y calcula los días de inventario como existencias divididas por la venta diaria.",
     "Un producto agotado es una venta perdida, y si es un producto estrella la pérdida es grande.",
     "Generar órdenes de compra priorizadas con el contacto del proveedor.",
     "Un libro ya está en cero unidades y el Smartwatch Zenit Plus, primero en ingresos, se agota en 22 días.",
     "Menos ventas perdidas y compras planificadas con anticipación."),
    ("C7. Tasa de Rechazo por Medio de Pago",
     "Por medio de pago: intentos, pagos aprobados, pagos rechazados, tasa de rechazo, monto aprobado y monto "
     "rechazado.",
     "Separa el array embebido de pagos con $unwind, agrupa por medio de pago usando $cond para contar y "
     "sumar según el estado del intento y calcula la tasa de rechazo.",
     "Cada rechazo es un cliente que puede abandonar la compra.",
     "Negociar con la pasarela de pagos y sugerir Nequi o PSE cuando falla la tarjeta.",
     "La tarjeta de crédito rechaza el 14,3 % de los intentos, unos 64 millones de pesos en compras que no "
     "pasaron al primer intento; Nequi solo rechaza el 7 %.",
     "Mayor conversión en el proceso de pago y menos carritos abandonados."),
]
for nombre, busca, pipeline, importancia, decision, hallazgo, impacto in consultas:
    titulo(nombre, 2)
    titulo4("Información que busca", busca)
    titulo4("Funcionamiento del pipeline", pipeline)
    if nombre.startswith("C1"):
        codigo([
            "db.pedidos.aggregate([",
            "  { $match: {",
            "      estado: { $in: [\"Pagado\", \"Enviado\", \"Entregado\"] },",
            "      fecha: { $gte: ISODate(\"2026-01-01\"),",
            "               $lt: ISODate(\"2027-01-01\") } } },",
            "  { $unwind: \"$items\" },",
            "  { $group: { _id: \"$items.producto_id\",",
            "              unidades: { $sum: \"$items.cantidad\" },",
            "              ingresos: { $sum: \"$items.subtotal\" },",
            "              pedidos: { $sum: 1 } } },",
            "  { $sort: { ingresos: -1 } },",
            "  { $limit: 10 },",
            "  { $lookup: { from: \"productos\", localField: \"_id\",",
            "               foreignField: \"_id\", as: \"producto\" } },",
            "  { $unwind: \"$producto\" },",
            "  { $lookup: { from: \"categorias\",",
            "               localField: \"producto.categoria_id\",",
            "               foreignField: \"_id\", as: \"categoria\" } },",
            "  { $unwind: \"$categoria\" },",
            "  { $lookup: { from: \"proveedores\",",
            "               localField: \"producto.proveedor_id\",",
            "               foreignField: \"_id\", as: \"proveedor\" } },",
            "  { $unwind: \"$proveedor\" },",
            "  { $project: { /* nombres y margen_bruto = ingresos",
            "                   menos unidades por costo */ } }",
            "])",
        ])
    titulo4("Importancia", importancia)
    titulo4("Decisión que apoya", decision)
    titulo4("Hallazgo", hallazgo)
    titulo4("Impacto para la organización", impacto)

# ================================================================= aplicación
titulo("Herramienta Complementaria: Aplicación de Escritorio", 1)
parrafo([("Además de la migración diseñada para TiendaOnline, el repositorio incluye una aplicación de "
          "escritorio desarrollada con PySide6 (carpeta ", {}), ("migrador", C),
         (") que migra cualquier base de datos de SQL Server hacia MongoDB. La aplicación lee el esquema "
          "desde las vistas de sistema (tablas, llaves primarias, llaves foráneas y restricciones únicas), "
          "mide cuántos registros hijos tiene cada registro padre y propone, para cada tabla, si debe ser una "
          "colección propia, un array embebido o un subdocumento, con la justificación de cada decisión.", {})])
parrafo("Las reglas que aplica son las mismas que se usaron para justificar el modelo de este proyecto: una "
        "tabla que otras referencian o que no depende de ninguna se mantiene como colección; una tabla que "
        "nadie referencia y que tiene un padre claro se embebe en él, como subdocumento si la relación es de "
        "uno a uno; una tabla que relaciona varias entidades sin un dueño claro se mantiene como colección con "
        "referencias; y una relación con más de 500 hijos por padre no se embebe para no acercarse al límite de "
        "16 MB por documento. El usuario puede cambiar cualquier decisión antes de migrar, y al final la "
        "aplicación verifica el resultado y genera un reporte. La Figura 3 muestra la pantalla del modelo.")
figura("Pantalla del Modelo Propuesto en la Aplicación de Escritorio", IMG / "figura_aplicacion.png", 6.5,
       "La aplicación propone para DetallePedido un array embebido en Pedidos, muestra la justificación y "
       "una vista previa del documento resultante.")
parrafo("Con TiendaOnline la aplicación propone siete colecciones, casi el mismo modelo diseñado en este "
        "proyecto. La única diferencia es Direcciones: la aplicación la mantiene como colección porque Pedidos "
        "la referencia, mientras que en el diseño manual se embebió en los clientes y se copió en cada pedido, "
        "una decisión de negocio que el esquema no revela. La verificación de la aplicación arrojó 39 de 39 "
        "controles correctos. La cifra es menor que la del verificador del proyecto (79) porque la aplicación "
        "construye controles genéricos a partir del esquema, mientras que el verificador del proyecto incluye "
        "reglas propias del negocio y cuenta cada campo de la muestra como un control.")

# ================================================================= presentación
titulo("Presentación de la Sustentación", 1)
parrafo([("La presentación utilizada en la sustentación se encuentra en ", {}),
         ("presentacion/Sustentacion_Taller3.pptx", C),
         (" y se genera con ", {}), ("presentacion/generar_presentacion.py", C),
         (". Tiene 22 diapositivas con notas para el expositor y sigue este recorrido: caso de estudio, modelo "
          "relacional, proceso de migración, modelo documental, decisiones de Embedding y Referencing, "
          "verificación, una diapositiva por cada consulta con su gráfico, hallazgo y decisión, la aplicación "
          "de escritorio y las conclusiones.", {})])

# ================================================================= conclusiones
titulo("Conclusiones", 1)
parrafo("El modelo documental se diseñó a partir de cómo se usan los datos y no de la estructura de las tablas. "
        "El pedido, que es la entidad central del negocio, se lee y se escribe completo en un solo documento, "
        "mientras que las entidades con existencia propia y que cambian con el tiempo se referencian.")
parrafo("La migración es verificable: 79 controles automáticos confirman que se conservaron las cantidades, "
        "los valores monetarios al centavo y todas las referencias, incluso en las tablas que se fusionaron "
        "dentro de otros documentos.")
parrafo("Las consultas demuestran que el modelo responde preguntas reales del negocio sobre inventario, "
        "proveedores, transportadoras, clientes y medios de pago, y que cada resultado apunta a una decisión "
        "concreta. Finalmente, la aplicación de escritorio muestra que los criterios de modelado aplicados son "
        "reglas explícitas que pueden reutilizarse con otras bases de datos.")

# ================================================================= apéndice
salto_pagina()
titulo("Apéndice", 1)
titulo("Instrucciones de Ejecución", 2)
parrafo("Los siguientes comandos se ejecutan desde la carpeta raíz del repositorio, en este orden. Requieren "
        "SQL Server con la instancia SQLEXPRESS, el controlador ODBC para SQL Server, MongoDB y Python.")
pasos = [
    ("Instalar las dependencias de Python:", "pip install -r requirements.txt"),
    ("Crear la base de datos y las tablas:", r"sqlcmd -S .\SQLEXPRESS -E -C -f 65001 -i sql\01_esquema.sql"),
    ("Cargar los datos:", r"sqlcmd -S .\SQLEXPRESS -E -C -f 65001 -i sql\02_carga_datos.sql"),
    ("Verificar la base relacional:", r"sqlcmd -S .\SQLEXPRESS -E -C -f 65001 -W -i sql\03_verificacion_sql.sql"),
    ("Migrar a MongoDB:", "python migracion/migrar.py"),
    ("Verificar la migración:", "python migracion/verificar.py"),
    ("Ejecutar las consultas:", "mongosh tienda_online --quiet consultas/consultas.js"),
    ("Abrir la aplicación de escritorio:", "python migrador_app.py"),
]
for i, (texto, comando) in enumerate(pasos, 1):
    parrafo(f"{i}. {texto}", sangria=False)
    codigo([comando])

doc.save(SALIDA)
print(f"Documento generado: {SALIDA}")
